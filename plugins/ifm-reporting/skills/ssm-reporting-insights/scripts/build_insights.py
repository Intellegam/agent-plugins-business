#!/usr/bin/env python3
"""
SSM Reporting Insights — the value-add layer on top of the monthly SSM quotes report.

Takes several monthly quotes files (chronological order) and produces one
workbook with:

  Tab "Trends"        : % active dealers and total quotes per brand/market per
                        month, with the change from first to last month and a
                        top-declines / top-improvements summary. Catches drops
                        like Hyundai Sweden 92% -> 58% (Jan-Jul 26) that are
                        invisible in single-month spot reports.
  Tab "Worst Markets" : campaign target list for the latest month — lowest
                        % active first (min network size filter), each with its
                        trend so "falling" markets can be rescued before they
                        reach the bottom.
  Tab "Data Quality"  : per input file: row counts and sanity checks, every
                        silently dropped dashboard-brand market (dealers +
                        quotes), and markets with >100% active (stale network
                        size). Makes the ~1,000 invisible quotes per month
                        visible.
  Tabs "Journey ..."  : (optional, --journey <Franchise>) dealer-level
                        month-by-month usage journey + movers list, as built
                        for Nissan Europe's webinar tracking.

Usage:
    python3 build_insights.py --ref <network_reference.csv> --out <out.xlsx>
        [--min-network 20] [--journey Nissan] <monthly files in date order...>
"""
import argparse
import collections

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from common import load_monthly, load_reference, category, CAT_RANK

ARIAL = "Arial"
HDRF = PatternFill("solid", fgColor="1F4E79")
F_RED = PatternFill("solid", fgColor="FFC7CE")
F_AMB = PatternFill("solid", fgColor="FFEB9C")
F_GRN = PatternFill("solid", fgColor="C6EFCE")
F_GRY = PatternFill("solid", fgColor="E7E6E6")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
CATF = {"NONE": F_GRY, "LOW": F_RED, "MEDIUM": F_AMB, "HIGH": F_GRN}


def hdr(ws, row, labels, widths=None):
    for i, h in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        c.fill = HDRF
        c.border = THIN
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def cell(ws, r, c, v, fmt=None, fill=None, bold=False, italic=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name=ARIAL, size=10, bold=bold, italic=italic)
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    x.border = THIN
    return x


def pct_fill(p):
    return F_RED if p < 0.6 else (F_AMB if p < 0.8 else F_GRN)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", help="monthly quotes files, oldest first")
    ap.add_argument("--ref", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--min-network", type=int, default=20,
                    help="minimum network size for the Worst Markets list")
    ap.add_argument("--journey", action="append", default=[],
                    help="add a dealer-journey tab for this franchise (repeatable)")
    args = ap.parse_args()

    ref_rows, lookup = load_reference(args.ref)
    months, frames, quality = [], {}, []
    for f in args.files:
        m, df = load_monthly(f)
        months.append(m)
        frames[m] = df

    # per market per month: active dealers + quotes (reference markets only)
    act = collections.defaultdict(dict)   # (brand,label) -> {month: n}
    qts = collections.defaultdict(dict)
    dropped = collections.defaultdict(dict)  # per month: (fr,co) -> (dealers, quotes)
    for m in months:
        df = frames[m]
        for (fr, co), grp in df.groupby(["Franchise", "Country"]):
            r = lookup.get((fr, co))
            if r is None:
                dash_franchises = {x["franchise"] for x in ref_rows}
                if fr in dash_franchises:
                    dropped[m][(fr, co)] = (len(grp), int(grp["Num Quotes"].sum()))
                continue
            key = (r["brand_label"], r["country_label"])
            act[key][m] = act[key].get(m, 0) + len(grp)
            qts[key][m] = qts[key].get(m, 0) + int(grp["Num Quotes"].sum())

    net = {(r["brand_label"], r["country_label"]): r["network"] for r in ref_rows}
    wb = Workbook()

    # ---------------- Overview (the one-look summary — keep it simple) ------
    ov = wb.active
    ov.title = "Overview"
    first, last = months[0], months[-1]
    total_net = sum(net.values())
    act_first = sum(act[k].get(first, 0) for k in net)
    act_last = sum(act[k].get(last, 0) for k in net)
    q_last = sum(qts[k].get(last, 0) for k in net)
    q_prev = sum(qts[k].get(months[-2], 0) for k in net) if len(months) > 1 else q_last
    kpi_deltas = []
    for k, n in net.items():
        if n >= args.min_network:
            kpi_deltas.append((k, (act[k].get(last, 0) - act[k].get(first, 0)) / n))
    falling = sorted([x for x in kpi_deltas if x[1] <= -0.05], key=lambda x: x[1])
    rising = sorted([x for x in kpi_deltas if x[1] >= 0.05], key=lambda x: -x[1])
    inv_dealers = sum(v[0] for v in dropped.get(last, {}).values())
    inv_quotes = sum(v[1] for v in dropped.get(last, {}).values())

    ov["A1"] = f"SSM Usage — What matters this month ({last})"
    ov["A1"].font = Font(name=ARIAL, bold=True, size=15)
    ov["A2"] = f"Period {first} – {last} · Menus app only · {len(net)} brand/markets tracked"
    ov["A2"].font = Font(name=ARIAL, size=10, color="808080")

    # KPI row
    kpis = [
        ("Network dealers", f"{total_net:,}", None),
        ("Active dealers", f"{act_last:,} ({act_last/total_net:.0%})",
         f"{(act_last-act_first)/total_net:+.0%} vs {first}"),
        (f"Quotes in {last}", f"{q_last:,}",
         f"{(q_last-q_prev)/q_prev:+.0%} vs prior month" if q_prev else ""),
        ("Markets falling ≥5pp", str(len(falling)), "need attention"),
        ("Usage invisible to report", f"{inv_quotes:,} quotes",
         f"{inv_dealers} dealers outside network list"),
    ]
    for i, (label, value, sub) in enumerate(kpis):
        col = 1 + i * 2
        c = ov.cell(row=4, column=col, value=label)
        c.font = Font(name=ARIAL, size=9, color="808080")
        c = ov.cell(row=5, column=col, value=value)
        c.font = Font(name=ARIAL, bold=True, size=13)
        if sub:
            c = ov.cell(row=6, column=col, value=sub)
            c.font = Font(name=ARIAL, size=9, color="808080")

    def block(start_row, title, rows_, fill, empty_text):
        cell(ov, start_row, 1, title, bold=True)
        r_ = start_row + 1
        if not rows_:
            cell(ov, r_, 1, empty_text, italic=True)
            return r_ + 1
        for k, d in rows_:
            a0 = act[k].get(first, 0) / net[k]
            a1 = act[k].get(last, 0) / net[k]
            cell(ov, r_, 1, f"{k[0]} {k[1]}")
            cell(ov, r_, 2, f"{a0:.0%} → {a1:.0%}", fill=fill)
            cell(ov, r_, 3, f"{d:+.0%}")
            r_ += 1
        return r_ + 1

    r0 = 8
    r0 = block(r0, f"ACT NOW — falling markets (biggest first, network ≥ {args.min_network}):",
               falling[:5], F_RED, "No market fell by 5pp or more. Good.")
    worst_now = sorted(((k, n) for k, n in net.items() if n >= args.min_network),
                       key=lambda kn: act[kn[0]].get(last, 0) / kn[1])[:5]
    cell(ov, r0, 1, "CAMPAIGN TARGETS — lowest activation now (see 'Worst Markets' tab for the full list):",
         bold=True)
    r0 += 1
    for k, n in worst_now:
        p = act[k].get(last, 0) / n
        cell(ov, r0, 1, f"{k[0]} {k[1]}")
        cell(ov, r0, 2, f"{p:.0%} of {n}", fill=pct_fill(p))
        cell(ov, r0, 3, f"{qts[k].get(last, 0):,} quotes")
        r0 += 1
    r0 += 1
    r0 = block(r0, "WHAT WORKS — biggest improvements (copy the playbook):",
               rising[:3], F_GRN, "No market improved by 5pp or more this period.")
    cell(ov, r0, 1, "Detail: Trends · Worst Markets · Data Quality tabs"
         + (" · Journey tabs per franchise" if args.journey else ""), italic=True)
    for col_letter, w in zip("ABCDEFGHIJ", [34, 16, 10, 16, 8, 18, 8, 20, 8, 22]):
        ov.column_dimensions[col_letter].width = w

    # ---------------- Trends ----------------
    ws = wb.create_sheet("Trends")
    ws["A1"] = f"SSM usage trends — {months[0]} to {months[-1]} (Menus only, % of network active)"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=13)
    labels = ["Brand", "Market", "Network"] + [f"{m} %act" for m in months] + \
             ["Δ %act", ] + [f"{m} quotes" for m in months] + ["Δ quotes"]
    hdr(ws, 3, labels, [14, 22, 9] + [9] * len(months) + [9] + [10] * len(months) + [10])
    r = 4
    deltas = []
    for key in sorted(net):
        n = net[key]
        cell(ws, r, 1, key[0])
        cell(ws, r, 2, key[1])
        cell(ws, r, 3, n, "#,##0")
        pcts = []
        for i, m in enumerate(months):
            p = (act[key].get(m, 0) / n) if n else 0
            pcts.append(p)
            cell(ws, r, 4 + i, p, "0%", pct_fill(p))
        d = pcts[-1] - pcts[0]
        cell(ws, r, 4 + len(months), d, "+0%;-0%;0%",
             F_RED if d <= -0.05 else (F_GRN if d >= 0.05 else None), bold=True)
        qs = [qts[key].get(m, 0) for m in months]
        for i, q in enumerate(qs):
            cell(ws, r, 5 + len(months) + i, q, "#,##0")
        cell(ws, r, 5 + 2 * len(months), qs[-1] - qs[0], "+#,##0;-#,##0;0", bold=True)
        if n >= args.min_network:
            deltas.append((key, pcts[0], pcts[-1], d))
        r += 1
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{get_column_letter(5 + 2 * len(months))}{r - 1}"

    # summary: top declines / improvements
    r += 2
    deltas.sort(key=lambda x: x[3])
    cell(ws, r, 1, f"Biggest declines (network ≥ {args.min_network}):", bold=True)
    for k, p0, p1, d in deltas[:5]:
        r += 1
        cell(ws, r, 1, f"{k[0]} {k[1]}")
        cell(ws, r, 2, f"{p0:.0%} → {p1:.0%} ({d:+.0%})", fill=F_RED)
    r += 2
    cell(ws, r, 1, "Biggest improvements:", bold=True)
    for k, p0, p1, d in deltas[-3:][::-1]:
        r += 1
        cell(ws, r, 1, f"{k[0]} {k[1]}")
        cell(ws, r, 2, f"{p0:.0%} → {p1:.0%} ({d:+.0%})", fill=F_GRN)

    # ---------------- Worst Markets ----------------
    ws = wb.create_sheet("Worst Markets")
    last = months[-1]
    ws["A1"] = (f"Campaign targets — lowest % active in {last} "
                f"(network ≥ {args.min_network}), with trend since {months[0]}")
    ws["A1"].font = Font(name=ARIAL, bold=True, size=13)
    hdr(ws, 3, ["Brand", "Market", "Network", f"Active {last}", "% active",
                f"Quotes {last}", f"Trend since {months[0]}", "Suggested angle"],
        [14, 22, 9, 10, 9, 10, 16, 42])
    worst = sorted(((k, n) for k, n in net.items() if n >= args.min_network),
                   key=lambda kn: act[kn[0]].get(last, 0) / kn[1])
    r = 4
    for k, n in worst[:25]:
        a = act[k].get(last, 0)
        p = a / n if n else 0
        d = p - ((act[k].get(months[0], 0) / n) if n else 0)
        trend = "falling" if d <= -0.05 else ("improving" if d >= 0.05 else "flat")
        angle = ("Win-back: usage is falling — contact before it hits zero" if trend == "falling"
                 else "Activation: market never picked up — onboarding/webinar candidate" if p < 0.3
                 else "Nudge: stable but below potential")
        cell(ws, r, 1, k[0]); cell(ws, r, 2, k[1]); cell(ws, r, 3, n, "#,##0")
        cell(ws, r, 4, a, "#,##0"); cell(ws, r, 5, p, "0%", pct_fill(p))
        cell(ws, r, 6, qts[k].get(last, 0), "#,##0")
        cell(ws, r, 7, f"{trend} ({d:+.0%})",
             fill=F_RED if trend == "falling" else (F_GRN if trend == "improving" else None))
        cell(ws, r, 8, angle)
        r += 1
    ws.freeze_panes = "A4"

    # ---------------- Data Quality ----------------
    ws = wb.create_sheet("Data Quality")
    ws["A1"] = "Data quality & silently dropped markets"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=13)
    hdr(ws, 3, ["Month", "Rows (Menus, EUR+MEA)", "In usual range (7,500-12,500)?",
                "Dropped dashboard-brand markets", "Dropped dealers", "Dropped quotes"],
        [12, 18, 22, 30, 12, 12])
    r = 4
    for m in months:
        nrows = len(frames[m])
        ok = 7500 <= nrows <= 12500
        dd = dropped.get(m, {})
        cell(ws, r, 1, m)
        cell(ws, r, 2, nrows, "#,##0")
        cell(ws, r, 3, "yes" if ok else "CHECK", None, None if ok else F_AMB)
        cell(ws, r, 4, len(dd), "#,##0", F_AMB if dd else None)
        cell(ws, r, 5, sum(v[0] for v in dd.values()), "#,##0")
        cell(ws, r, 6, sum(v[1] for v in dd.values()), "#,##0")
        r += 1
    r += 1
    cell(ws, r, 1, "Dropped markets in detail (dashboard brands only — decide with Rob "
                   "whether to add them to the network reference):", bold=True)
    r += 1
    hdr(ws, r, ["Month", "Franchise", "Country", "Active dealers", "Quotes"], None)
    r += 1
    for m in months:
        for (fr, co), (dl, q) in sorted(dropped.get(m, {}).items()):
            cell(ws, r, 1, m); cell(ws, r, 2, fr); cell(ws, r, 3, co)
            cell(ws, r, 4, dl, "#,##0"); cell(ws, r, 5, q, "#,##0")
            r += 1
    r += 1
    cell(ws, r, 1, "Markets with >100% active (network size stale — refresh via dealer listing):",
         bold=True)
    r += 1
    for k, n in sorted(net.items()):
        a = act[k].get(last, 0)
        if n and a > n:
            cell(ws, r, 1, f"{k[0]} {k[1]}: {a} active vs network {n}", fill=F_AMB)
            r += 1

    # ---------------- optional Journey tabs ----------------
    for franchise in args.journey:
        per = {}
        for m in months:
            df = frames[m]
            sel = df[df["Franchise"] == franchise]
            if "Dealer Name" not in sel.columns:
                continue
            per[m] = sel.groupby(["Country", "Dealer Name"])["Num Quotes"].sum()
        dealers = sorted(set().union(*[set(s.index) for s in per.values()])) if per else []
        ws = wb.create_sheet(f"Journey {franchise}"[:31])
        hdr(ws, 1, ["Country", "Dealer Name"] + months + ["Trend"],
            [18, 42] + [8] * len(months) + [12])
        r = 2
        for (co, dn) in dealers:
            vals = [int(per[m].get((co, dn), 0)) for m in months]
            d = CAT_RANK[category(vals[-1])] - CAT_RANK[category(vals[0])]
            cell(ws, r, 1, co)
            cell(ws, r, 2, dn)
            for i, v in enumerate(vals):
                cell(ws, r, 3 + i, v, "#,##0", CATF[category(v)])
            cell(ws, r, 3 + len(months),
                 "↑ improving" if d > 0 else ("↓ declining" if d < 0 else "→ stable"))
            r += 1
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(3 + len(months))}{r - 1}"

    wb.save(args.out)
    print(f"Saved: {args.out}")
    print(f"Months: {', '.join(months)} | markets tracked: {len(net)} | "
          f"dropped markets flagged: {sum(len(v) for v in dropped.values())}")


if __name__ == "__main__":
    main()
